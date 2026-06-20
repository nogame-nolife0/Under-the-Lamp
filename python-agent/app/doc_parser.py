"""Word 文档解析：提取文本、公式（OMML）与嵌入图片。"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

from app.latex_normalize import normalize_latex_delimiters

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/officeDocument/2006/main"


@dataclass
class DocBlock:
    kind: str  # text | image | break
    content: str = ""


def _local_tag(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def _omml_to_text(elem: ET.Element) -> str:
    tag = _local_tag(elem.tag)

    if tag == "t":
        return elem.text or ""

    if tag == "r":
        return "".join(_omml_to_text(child) for child in elem)

    if tag == "e":
        return "".join(_omml_to_text(child) for child in elem)

    if tag == "acc":
        accent_chr = "̅"
        chr_el = elem.find(qn("m:chr"))
        if chr_el is not None:
            accent_chr = chr_el.get(qn("m:val"), accent_chr) or accent_chr
        base = elem.find(qn("m:e"))
        inner = _omml_to_text(base) if base is not None else ""
        if accent_chr in {"̅", "¯", "-", "‾", "^"}:
            return f"\\overline{{{inner}}}"
        return f"{inner}{accent_chr}"

    if tag == "bar":
        base = elem.find(qn("m:e"))
        inner = _omml_to_text(base) if base is not None else ""
        return f"\\overline{{{inner}}}"

    if tag == "sup":
        base = elem.find(qn("m:e"))
        sup = elem.find(qn("m:sup"))
        b = _omml_to_text(base) if base is not None else ""
        s = _omml_to_text(sup) if sup is not None else ""
        return f"{b}^{{{s}}}"

    if tag == "sub":
        base = elem.find(qn("m:e"))
        sub = elem.find(qn("m:sub"))
        b = _omml_to_text(base) if base is not None else ""
        s = _omml_to_text(sub) if sub is not None else ""
        return f"{b}_{{{s}}}"

    if tag == "sSub":
        base = elem.find(qn("m:e"))
        sub = elem.find(qn("m:sub"))
        b = _omml_to_text(base) if base is not None else ""
        s = _omml_to_text(sub) if sub is not None else ""
        return f"{b}_{{{s}}}"

    if tag == "sSup":
        base = elem.find(qn("m:e"))
        sup = elem.find(qn("m:sup"))
        b = _omml_to_text(base) if base is not None else ""
        s = _omml_to_text(sup) if sup is not None else ""
        return f"{b}^{{{s}}}"

    if tag == "f":
        num = elem.find(qn("m:num"))
        den = elem.find(qn("m:den"))
        n = _omml_to_text(num) if num is not None else ""
        d = _omml_to_text(den) if den is not None else ""
        return f"({n})/({d})"

    if tag == "d":
        beg = elem.find(qn("m:begChr"))
        end = elem.find(qn("m:endChr"))
        left = beg.get(qn("m:val"), "(") if beg is not None else "("
        right = end.get(qn("m:val"), ")") if end is not None else ")"
        e = elem.find(qn("m:e"))
        inner = _omml_to_text(e) if e is not None else ""
        return f"{left}{inner}{right}"

    if tag == "rad":
        base = elem.find(qn("m:e"))
        inner = _omml_to_text(base) if base is not None else ""
        return f"sqrt({inner})"

    if tag in {"oMath", "oMathPara"}:
        return "".join(_omml_to_text(child) for child in elem)

    return "".join(_omml_to_text(child) for child in elem)


def _walk_run(run_elem: ET.Element) -> str:
    parts: list[str] = []
    for child in run_elem:
        tag = _local_tag(child.tag)
        if tag == "t":
            parts.append(child.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag in {"oMath", "oMathPara"}:
            parts.append(_omml_to_text(child))
    return "".join(parts)


def _walk_container(elem: ET.Element) -> str:
    parts: list[str] = []
    for child in elem:
        tag = _local_tag(child.tag)
        if tag == "r":
            parts.append(_walk_run(child))
        elif tag in {"oMath", "oMathPara"}:
            parts.append(_omml_to_text(child))
        elif tag == "hyperlink":
            parts.append(_walk_container(child))
    return "".join(parts)


def _paragraph_to_text(paragraph) -> str:
    text = _walk_container(paragraph._element).strip()
    if text:
        return text
    return paragraph.text.strip()


def _content_type_to_ext(content_type: str) -> str:
    mapping = {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/gif": ".gif",
        "image/bmp": ".bmp",
        "image/tiff": ".tif",
        "image/x-emf": ".emf",
        "image/emf": ".emf",
        "image/x-wmf": ".wmf",
        "image/wmf": ".wmf",
    }
    return mapping.get(content_type, ".png")


def _save_image_part(image_part, out_dir: Path, counter: int) -> Path:
    ext = _content_type_to_ext(getattr(image_part, "content_type", "image/png"))
    file_path = out_dir / f"img_{counter:03d}_{uuid.uuid4().hex[:8]}{ext}"
    file_path.write_bytes(image_part.blob)
    return file_path


def _iter_elements_by_tag(elem: ET.Element, local_names: set[str]):
    if _local_tag(elem.tag) in local_names:
        yield elem
    for child in elem:
        yield from _iter_elements_by_tag(child, local_names)


def _collect_images_from_element(elem, doc, out_dir: Path, seen: set[str], counter: list[int]) -> list[Path]:
    paths: list[Path] = []

    for blip in _iter_elements_by_tag(elem, {"blip"}):
        embed = blip.get(qn("r:embed"))
        if not embed or embed in seen:
            continue
        seen.add(embed)
        try:
            image_part = doc.part.related_parts[embed]
            counter[0] += 1
            paths.append(_save_image_part(image_part, out_dir, counter[0]))
        except KeyError:
            continue

    for imagedata in _iter_elements_by_tag(elem, {"imagedata"}):
        embed = imagedata.get(qn("r:id"))
        if not embed or embed in seen:
            continue
        seen.add(embed)
        try:
            image_part = doc.part.related_parts[embed]
            counter[0] += 1
            paths.append(_save_image_part(image_part, out_dir, counter[0]))
        except KeyError:
            continue

    return paths


def _normalize_text_block(text: str) -> str:
    return normalize_latex_delimiters(text.strip()) if text else text


def _element_has_embedded_image(elem: ET.Element) -> bool:
    for node in _iter_elements_by_tag(elem, {"blip", "imagedata", "drawing", "pict", "object"}):
        if _local_tag(node.tag) in {"blip", "imagedata"}:
            return True
        if _local_tag(node.tag) in {"drawing", "pict", "object"}:
            return True
    return False


def _run_to_blocks(run_elem: ET.Element, doc, out_dir: Path, seen: set[str], counter: list[int]) -> list[DocBlock]:
    blocks: list[DocBlock] = []
    text_buf: list[str] = []

    def flush_text() -> None:
        if not text_buf:
            return
        normalized = _normalize_text_block("".join(text_buf))
        text_buf.clear()
        if normalized:
            blocks.append(DocBlock("text", normalized))

    for child in run_elem:
        tag = _local_tag(child.tag)
        if tag == "t":
            text_buf.append(child.text or "")
        elif tag == "tab":
            text_buf.append("\t")
        elif tag in {"br", "cr"}:
            text_buf.append(" ")
        elif tag in {"oMath", "oMathPara"}:
            flush_text()
            normalized = _normalize_text_block(_omml_to_text(child))
            if normalized:
                blocks.append(DocBlock("text", normalized))
        elif _element_has_embedded_image(child):
            flush_text()
            for image_path in _collect_images_from_element(child, doc, out_dir, seen, counter):
                blocks.append(DocBlock("image", str(image_path)))
        elif tag in {"hyperlink", "sdt", "ins", "smartTag"}:
            flush_text()
            if _element_has_embedded_image(child):
                for image_path in _collect_images_from_element(child, doc, out_dir, seen, counter):
                    blocks.append(DocBlock("image", str(image_path)))
            nested_text = _walk_container(child)
            if nested_text:
                text_buf.append(nested_text)
        elif tag == "r":
            flush_text()
            blocks.extend(_run_to_blocks(child, doc, out_dir, seen, counter))
        else:
            if _element_has_embedded_image(child):
                flush_text()
                for image_path in _collect_images_from_element(child, doc, out_dir, seen, counter):
                    blocks.append(DocBlock("image", str(image_path)))
            else:
                nested_text = _walk_run(child) if _local_tag(child.tag) == "r" else _walk_container(child)
                if nested_text:
                    text_buf.append(nested_text)

    flush_text()
    return blocks


def _paragraph_to_blocks(paragraph, doc, out_dir: Path, seen: set[str], counter: list[int]) -> list[DocBlock]:
    blocks: list[DocBlock] = []
    element = paragraph._element

    def walk(elem: ET.Element) -> None:
        tag = _local_tag(elem.tag)
        if tag == "r":
            blocks.extend(_run_to_blocks(elem, doc, out_dir, seen, counter))
            return
        if tag in {"oMath", "oMathPara"}:
            normalized = _normalize_text_block(_omml_to_text(elem))
            if normalized:
                blocks.append(DocBlock("text", normalized))
            return
        if tag in {"hyperlink", "ins", "sdt", "smartTag"}:
            for child in elem:
                walk(child)
            return
        for child in elem:
            walk(child)

    for child in element:
        walk(child)

    if not blocks:
        text = _normalize_text_block(_paragraph_to_text(paragraph))
        if text:
            blocks.append(DocBlock("text", text))
    if blocks:
        blocks.append(DocBlock("break"))
    return blocks


def _cell_to_blocks(cell, doc, out_dir: Path, seen: set[str], counter: list[int]) -> list[DocBlock]:
    blocks: list[DocBlock] = []
    for paragraph in cell.paragraphs:
        blocks.extend(_paragraph_to_blocks(paragraph, doc, out_dir, seen, counter))
    return blocks


def extract_document_blocks(file_path: str) -> list[DocBlock]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Word 文件不存在: {file_path}")
    if path.suffix.lower() not in {".docx"}:
        raise ValueError("仅支持 .docx 格式，请将 .doc 另存为 .docx")

    doc = Document(str(path))
    out_dir = path.parent / f".extracted_images_{path.stem}"
    out_dir.mkdir(parents=True, exist_ok=True)

    blocks: list[DocBlock] = []
    seen: set[str] = set()
    counter = [0]

    for paragraph in doc.paragraphs:
        blocks.extend(_paragraph_to_blocks(paragraph, doc, out_dir, seen, counter))

    for table in doc.tables:
        for row in table.rows:
            row_blocks: list[DocBlock] = []
            for cell in row.cells:
                row_blocks.extend(_cell_to_blocks(cell, doc, out_dir, seen, counter))
            if row_blocks:
                blocks.extend(row_blocks)
                blocks.append(DocBlock("break"))

    if not blocks:
        raise ValueError("Word 文件中未提取到有效内容")

    return blocks


def blocks_to_inline_text(blocks: list[DocBlock]) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.kind == "text":
            parts.append(block.content)
        elif block.kind == "break":
            parts.append("\n")
    return "".join(parts)


def read_docx_text(file_path: str) -> str:
    blocks = extract_document_blocks(file_path)
    text = blocks_to_inline_text(blocks)
    if not text.strip():
        image_count = sum(1 for block in blocks if block.kind == "image")
        if image_count:
            return f"[文档仅含 {image_count} 张嵌入图片，需视觉识别]"
        raise ValueError("Word 文件中未提取到有效文本")
    return text


def count_images(blocks: list[DocBlock]) -> int:
    return sum(1 for block in blocks if block.kind == "image")


def truncate_text(text: str, max_chars: int = 80000) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[文本已截断，后续内容未送入模型]"
